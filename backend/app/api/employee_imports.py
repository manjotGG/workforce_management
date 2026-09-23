import csv
import io
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from fastapi import APIRouter, UploadFile, File, Depends, HTTPException, status
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.core.security import require_roles
from app.models import Employee, Department, Shift, UserRole, User

router = APIRouter(prefix="/api/employees", tags=["employees"])


import csv
import io
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any
from fastapi import APIRouter, UploadFile, File, Depends, HTTPException, status
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.core.security import require_roles
from app.models import Employee, Department, Shift, UserRole, User

router = APIRouter(prefix="/api/employees", tags=["employees"])


def _parse_date(val: Any) -> date | None:
    if val is None or val == "":
        return None
    if isinstance(val, date):
        return val
    if isinstance(val, datetime):
        return val.date()
    val_str = str(val).strip()
    if not val_str:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y %H:%M:%S"):
        try:
            return datetime.strptime(val_str, fmt).date()
        except Exception:
            continue
    try:
        return datetime.fromisoformat(val_str).date()
    except Exception:
        return None


def _clean_str(val: Any) -> str | None:
    if val is None or val == "":
        return None
    if isinstance(val, float) and val.is_integer():
        return str(int(val))
    val_s = str(val).strip()
    return val_s if val_s else None


def _resolve_department(db: Session, val: Any) -> int | None:
    cleaned = _clean_str(val)
    if not cleaned:
        return None
    if cleaned.isdigit():
        d = db.get(Department, int(cleaned))
        return d.id if d else None
    d = db.query(Department).filter(Department.name.ilike(cleaned)).first()
    return d.id if d else None


def _resolve_shift(db: Session, val: Any) -> int | None:
    cleaned = _clean_str(val)
    if not cleaned:
        return None
    if cleaned.isdigit():
        s = db.get(Shift, int(cleaned))
        return s.id if s else None
    s = db.query(Shift).filter(Shift.name.ilike(cleaned)).first()
    return s.id if s else None


def _get_val_by_aliases(raw_dict: dict, aliases: list[str]) -> Any:
    norm_map = {}
    for k, v in raw_dict.items():
        if k is not None:
            norm_k = str(k).strip().lower().replace(" ", "_").replace("-", "_")
            norm_map[norm_k] = v
    for alias in aliases:
        norm_alias = alias.strip().lower().replace(" ", "_").replace("-", "_")
        if norm_alias in norm_map and norm_map[norm_alias] not in (None, ""):
            return norm_map[norm_alias]
    return None


def _rows_from_csv_bytes(b: bytes) -> list[dict]:
    text = b.decode(errors="replace")
    f = io.StringIO(text)
    reader = csv.DictReader(f)
    rows = []
    for row in reader:
        rows.append({k.strip() if k else "": (v.strip() if isinstance(v, str) else v) for k, v in row.items()})
    return rows


def _rows_from_xlsx_fileobj(fileobj) -> list[dict]:
    try:
        import openpyxl
    except ImportError:
        raise RuntimeError("openpyxl is required to parse xlsx files. Please install openpyxl.")

    wb = openpyxl.load_workbook(fileobj, read_only=True, data_only=True)
    ws = wb.active or wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)

    headers = None
    for r in rows:
        if r and any(c is not None and str(c).strip() for c in r):
            headers = [str(h).strip() if h is not None else "" for h in r]
            break

    if not headers:
        return []

    result = []
    for r in rows:
        if not r or not any(c is not None and str(c).strip() for c in r):
            continue
        d = {}
        for k, v in zip(headers, r):
            if k:
                d[k] = v
        result.append(d)
    return result


@router.post("/import", status_code=status.HTTP_200_OK)
async def import_employees(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles(UserRole.ADMIN, UserRole.ATTENDANCE_OPERATOR)),
):
    """Import employees from CSV or XLSX. Returns per-row success/failure summary."""
    filename = (file.filename or "").lower()
    try:
        contents = await file.read()
        if filename.endswith(".csv"):
            rows = _rows_from_csv_bytes(contents)
        elif filename.endswith(('.xls', '.xlsx')):
            fh = io.BytesIO(contents)
            rows = _rows_from_xlsx_fileobj(fh)
        else:
            raise HTTPException(status_code=400, detail="Unsupported file type. Upload CSV or XLSX.")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to read file: {str(e)}")

    total = 0
    success = 0
    failed = []

    for idx, raw in enumerate(rows, start=1):
        total += 1
        try:
            emp_id_val = _get_val_by_aliases(raw, ['employee_id', 'emp_id', 'employee_code', 'emp_code', 'id'])
            name_val = _get_val_by_aliases(raw, ['name', 'full_name', 'employee_name'])

            emp_id = _clean_str(emp_id_val)
            name = _clean_str(name_val)

            if not emp_id:
                raise ValueError('employee_id is required')
            if not name:
                raise ValueError('name is required')

            exists = db.query(Employee).filter(Employee.employee_id == emp_id).first()
            if exists:
                raise ValueError(f'Employee with ID "{emp_id}" already exists')

            dept_val = _get_val_by_aliases(raw, ['department', 'dept', 'department_name', 'department_id'])
            shift_val = _get_val_by_aliases(raw, ['shift', 'shift_name', 'shift_id', 'schedule'])
            phone_val = _get_val_by_aliases(raw, ['phone', 'mobile', 'phone_number', 'contact'])
            desig_val = _get_val_by_aliases(raw, ['designation', 'role', 'title', 'position'])
            salary_val = _get_val_by_aliases(raw, ['monthly_salary', 'salary', 'base_salary', 'pay'])
            joining_val = _get_val_by_aliases(raw, ['joining_date', 'date_of_joining', 'joining', 'doj', 'start_date'])

            dept_id = _resolve_department(db, dept_val)
            shift_id = _resolve_shift(db, shift_val)

            monthly_salary = None
            if salary_val not in (None, ""):
                try:
                    monthly_salary = Decimal(str(salary_val))
                except (InvalidOperation, TypeError):
                    raise ValueError('invalid monthly_salary')

            joining = _parse_date(joining_val)

            emp = Employee(
                employee_id=emp_id,
                name=name,
                phone=_clean_str(phone_val),
                department_id=dept_id,
                designation=_clean_str(desig_val),
                shift_id=shift_id,
                monthly_salary=monthly_salary,
                joining_date=joining,
                is_active=True,
            )

            db.add(emp)
            db.commit()
            db.refresh(emp)
            success += 1
        except Exception as e:
            db.rollback()
            failed.append({"row": idx, "error": str(e), "data": {str(k): str(v) for k, v in raw.items()}})

    return {"total_rows": total, "successful_rows": success, "failed_rows": len(failed), "failures": failed}

