import csv
import io
from datetime import datetime, date
from typing import Any
from fastapi import APIRouter, UploadFile, File, Depends, HTTPException, status
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.core.security import require_roles
from app.models import (
    AttendanceImport,
    RawAttendanceRecord,
    Attendance,
    Employee,
    User,
    UserRole,
)

router = APIRouter(prefix="/api/attendance", tags=["attendance"])


def _rows_from_csv_bytes(b: bytes) -> list[dict]:
    text = b.decode(errors="replace")
    f = io.StringIO(text)
    reader = csv.DictReader(f)
    rows = []
    for row in reader:
        rows.append({k.strip() if k else "": (v.strip() if isinstance(v, str) else v) for k, v in row.items()})
    return rows


def _rows_from_xlsx_fileobj(fileobj) -> list[dict]:
    try:
        import openpyxl
    except ImportError:
        raise RuntimeError("openpyxl is required to parse xlsx files. Please install openpyxl.")

    wb = openpyxl.load_workbook(fileobj, read_only=True, data_only=True)
    ws = wb.active or wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)

    headers = None
    for r in rows:
        if r and any(c is not None and str(c).strip() for c in r):
            headers = [str(h).strip() if h is not None else "" for h in r]
            break

    if not headers:
        return []

    result = []
    for r in rows:
        if not r or not any(c is not None and str(c).strip() for c in r):
            continue
        d = {}
        for k, v in zip(headers, r):
            if k:
                d[k] = v
        result.append(d)
    return result


def _get_val_by_aliases(raw_dict: dict, aliases: list[str]) -> Any:
    norm_map = {}
    for k, v in raw_dict.items():
        if k is not None:
            norm_k = str(k).strip().lower().replace(" ", "_").replace("-", "_")
            norm_map[norm_k] = v
    for alias in aliases:
        norm_alias = alias.strip().lower().replace(" ", "_").replace("-", "_")
        if norm_alias in norm_map and norm_map[norm_alias] not in (None, ""):
            return norm_map[norm_alias]
    return None


@router.post("/import_file", status_code=status.HTTP_200_OK)
async def import_attendance_file(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles(UserRole.ADMIN, UserRole.ATTENDANCE_OPERATOR)),
):
    """Upload and process attendance CSV/XLSX file. Returns per-row summary."""
    filename = (file.filename or "").lower()
    try:
        contents = await file.read()
        if filename.endswith(".csv"):
            rows = _rows_from_csv_bytes(contents)
        elif filename.endswith((_".xls", ".xlsx")):
            fh = io.BytesIO(contents)
            rows = _rows_from_xlsx_fileobj(fh)
        else:
            raise HTTPException(status_code=400, detail="Unsupported file type. Upload CSV or XLSX.")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to read file: {str(e)}")

    # create attendance import record
    att_import = AttendanceImport(original_filename=file.filename or "uploaded", file_type=filename.split('.')[-1])
    db.add(att_import)
    db.commit()
    db.refresh(att_import)

    total = 0
    success = 0
    failed = []
    duplicates = 0

    for idx, raw in enumerate(rows, start=1):
        total += 1
        try:
            emp_code = _get_val_by_aliases(raw, ["employee_id", "emp_id", "employee_code", "emp_code", "id"])
            date_val = _get_val_by_aliases(raw, ["attendance_date", "date", "day"])
            first_in_val = _get_val_by_aliases(raw, ["first_in", "in_time", "first"])
            last_out_val = _get_val_by_aliases(raw, ["last_out", "out_time", "last"])
            punch_type = _get_val_by_aliases(raw, ["punch_type", "type"])

            if not emp_code:
                raise ValueError("employee_id is required")

            # resolve employee by employee_id code
            emp = db.query(Employee).filter(Employee.employee_id == str(emp_code).strip()).first()
            emp_id = emp.id if emp else None

            # parse date
            att_date = None
            if date_val is not None and date_val != "":
                if isinstance(date_val, date):
                    att_date = date_val
                else:
                    try:
                        att_date = datetime.fromisoformat(str(date_val)).date()
                    except Exception:
                        try:
                            att_date = datetime.strptime(str(date_val), "%d/%m/%Y").date()
                        except Exception:
                            try:
                                att_date = datetime.strptime(str(date_val), "%Y-%m-%d").date()
                            except Exception:
                                raise ValueError("Invalid date format")
            else:
                raise ValueError("attendance_date is required")

            # parse times
            first_in_dt = None
            last_out_dt = None
            if first_in_val:
                # Handle several possible formats: datetime, time string, or Excel numeric
                if isinstance(first_in_val, (int, float)):
                    try:
                        # openpyxl provides from_excel for numeric serials
                        from openpyxl.utils.datetime import from_excel

                        first_in_dt = from_excel(first_in_val)
                    except Exception:
                        first_in_dt = None
                else:
                    try:
                        first_in_dt = datetime.fromisoformat(str(first_in_val))
                    except Exception:
                        try:
                            first_in_dt = datetime.strptime(str(first_in_val), "%H:%M:%S")
                        except Exception:
                            try:
                                first_in_dt = datetime.strptime(str(first_in_val), "%H:%M")
                            except Exception:
                                first_in_dt = None
            if last_out_val:
                if isinstance(last_out_val, (int, float)):
                    try:
                        from openpyxl.utils.datetime import from_excel

                        last_out_dt = from_excel(last_out_val)
                    except Exception:
                        last_out_dt = None
                else:
                    try:
                        last_out_dt = datetime.fromisoformat(str(last_out_val))
                    except Exception:
                        try:
                            last_out_dt = datetime.strptime(str(last_out_val), "%H:%M:%S")
                        except Exception:
                            try:
                                last_out_dt = datetime.strptime(str(last_out_val), "%H:%M")
                            except Exception:
                                last_out_dt = None

            # If parsed times are date-less (year 1900), attach the attendance_date for correct datetimes
            if first_in_dt is not None and first_in_dt.year == 1900 and att_date is not None:
                first_in_dt = datetime.combine(att_date, first_in_dt.time())
            if last_out_dt is not None and last_out_dt.year == 1900 and att_date is not None:
                last_out_dt = datetime.combine(att_date, last_out_dt.time())

            # If last_out earlier than first_in, assume punch crossed midnight and add one day
            if first_in_dt and last_out_dt and last_out_dt < first_in_dt:
                from datetime import timedelta

                last_out_dt = last_out_dt + timedelta(days=1)

            # compute worked minutes when possible
            computed_worked_minutes = None
            if first_in_dt and last_out_dt:
                diff = last_out_dt - first_in_dt
                computed_worked_minutes = int(diff.total_seconds() // 60)

            # prevent duplicate attendance for same emp/date
            if emp_id:
                existing = (
                    db.query(Attendance)
                    .filter(Attendance.employee_id == emp_id, Attendance.attendance_date == att_date)
                    .first()
                )
                if existing:
                    duplicates += 1
                    # still store raw record
                    raw_obj = RawAttendanceRecord(
                        attendance_import_id=att_import.id,
                        employee_id=emp_id,
                        device_employee_identifier=str(emp.employee_id),
                        punch_timestamp=first_in_dt or (last_out_dt or datetime.utcnow()),
                        punch_type=punch_type or "",
                        source="file",
                        raw_data=raw,
                    )
                    db.add(raw_obj)
                    db.commit()
                    continue

            # if employee not found, mark failed but still store raw
            if not emp_id:
                raw_obj = RawAttendanceRecord(
                    attendance_import_id=att_import.id,
                    employee_id=None,
                    device_employee_identifier=str(emp_code),
                    punch_timestamp=first_in_dt or (last_out_dt or datetime.utcnow()),
                    punch_type=punch_type or "",
                    source="file",
                    raw_data=raw,
                )
                db.add(raw_obj)
                db.commit()
                failed.append({"row": idx, "error": "Employee not found", "data": raw})
                continue

            # create or update attendance
            existing_att = (
                db.query(Attendance)
                .filter(Attendance.employee_id == emp_id, Attendance.attendance_date == att_date)
                .first()
            )
            status = "PRESENT"
            worked_minutes = None
            if existing_att:
                existing_att.first_in = existing_att.first_in or first_in_dt
                existing_att.last_out = last_out_dt or existing_att.last_out
                # set worked_minutes if computed
                if computed_worked_minutes is not None:
                    existing_att.worked_minutes = computed_worked_minutes
                db.add(existing_att)
                db.commit()
            else:
                new_att = Attendance(
                    employee_id=emp_id,
                    attendance_date=att_date,
                    status=status,
                    first_in=first_in_dt,
                    last_out=last_out_dt,
                    worked_minutes=computed_worked_minutes,
                    attendance_import_id=att_import.id,
                )
                db.add(new_att)
                db.commit()

            # store raw record
            raw_obj = RawAttendanceRecord(
                attendance_import_id=att_import.id,
                employee_id=emp_id,
                device_employee_identifier=str(emp.employee_id),
                punch_timestamp=first_in_dt or (last_out_dt or datetime.utcnow()),
                punch_type=punch_type or "",
                source="file",
                raw_data=raw,
            )
            db.add(raw_obj)
            db.commit()

            success += 1
        except Exception as e:
            db.rollback()
            failed.append({"row": idx, "error": str(e), "data": raw})

    # update import record
    att_import.total_rows = total
    att_import.successful_rows = success
    att_import.failed_rows = len(failed)
    att_import.duplicate_rows = duplicates
    att_import.processing_completed_at = datetime.utcnow()
    att_import.status = "COMPLETED" if len(failed) == 0 else ("PARTIAL" if success > 0 else "FAILED")
    att_import.error_summary = {"failures": failed[:50]} if failed else None
    db.add(att_import)
    db.commit()
    db.refresh(att_import)

    return {
        "total_rows": total,
        "successful_rows": success,
        "failed_rows": len(failed),
        "duplicate_rows": duplicates,
        "failures": failed,
        "import_id": att_import.id,
    }
